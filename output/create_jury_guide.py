from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(r'C:\Users\Lenovo\Desktop\ProofForge_Cyber_Jury_Guide.docx')
doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
sec.header_distance=Inches(.32); sec.footer_distance=Inches(.32)
BLACK='000000'; INK='202824'; GREEN='24563E'; MUTED='5F6D65'; NAVY='20352B'; PALE='EEF3EE'; STRIPE='F6F7F4'; BORDER='D6DED7'; WHITE='FFFFFF'
normal=doc.styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(INK); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.12
for nm,sz in [('Title',30),('Subtitle',14),('Heading 1',19),('Heading 2',14),('Heading 3',11.5)]:
 s=doc.styles[nm]; s.font.name='Aptos Display' if nm in ('Title','Heading 1') else 'Aptos'; s.font.size=Pt(sz); s.font.color.rgb=RGBColor.from_string(BLACK); s.font.bold=nm!='Subtitle'; s.paragraph_format.keep_with_next=True
 if nm=='Heading 1': s.paragraph_format.space_before=Pt(17); s.paragraph_format.space_after=Pt(7)
 if nm=='Heading 2': s.paragraph_format.space_before=Pt(12); s.paragraph_format.space_after=Pt(4)
# Remove the built-in blue Title rule so the cover remains clean and plain.
title_ppr=doc.styles['Title']._element.find(qn('w:pPr'))
if title_ppr is not None:
 pborder=title_ppr.find(qn('w:pBdr'))
 if pborder is not None: title_ppr.remove(pborder)
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('FREQUENCY  |  PROOFFORGE CYBER'); r.font.size=Pt(8); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(MUTED)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Team HM26-C3B5  ·  Jury preparation guide  ·  '); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
def shade(c,fill):
 p=c._tc.get_or_add_tcPr(); x=OxmlElement('w:shd'); x.set(qn('w:fill'),fill); p.append(x)
def borders(c):
 p=c._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders'); p.append(b)
 for e in ('top','left','bottom','right','insideH','insideV'):
  x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'5'); x.set(qn('w:color'),BORDER); b.append(x)
def table(head,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(head)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths): c.width=Inches(w)
 for i,v in enumerate(head):
  c=t.rows[0].cells[i]; c.text=v; shade(c,NAVY); borders(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for p in c.paragraphs:
   for r in p.runs: r.font.size=Pt(9.5); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(WHITE)
 for n,row in enumerate(rows):
  cells=t.add_row().cells
  for i,v in enumerate(row):
   c=cells[i]; c.text=v; borders(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if n%2: shade(c,STRIPE)
   if widths: c.width=Inches(widths[i])
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.05
    for r in p.runs: r.font.size=Pt(9.2)
 doc.add_paragraph().paragraph_format.space_after=Pt(1)
def p(text='',style=None): return doc.add_paragraph(text,style=style)
def bullet(text):
 x=doc.add_paragraph(style='List Bullet'); x.paragraph_format.space_after=Pt(3); x.add_run(text)
def num(text):
 x=doc.add_paragraph(style='List Number'); x.paragraph_format.space_after=Pt(4); x.add_run(text)
def h1(text): doc.add_heading(text,1)
def h2(text): doc.add_heading(text,2)
def h3(text): doc.add_heading(text,3)
def qa(question,answer):
 h3(question); x=doc.add_paragraph(); x.paragraph_format.space_after=Pt(6); r=x.add_run('Suggested answer  '); r.bold=True; r.font.color.rgb=RGBColor.from_string(GREEN); x.add_run(answer)

# Cover and opening context
x=doc.add_paragraph(style='Title'); x.paragraph_format.space_before=Pt(82); x.paragraph_format.space_after=Pt(8); x.add_run('ProofForge Cyber Jury Guide')
x=doc.add_paragraph(style='Subtitle'); x.paragraph_format.space_after=Pt(20); x.add_run('Project explanation, professional pitch, demo plan, jury questions and roadmap')
x=p('FREQUENCY  ·  Team HM26-C3B5'); x.runs[0].font.bold=True; x.runs[0].font.size=Pt(12); x.runs[0].font.color.rgb=RGBColor.from_string(GREEN)
p('A practical speaking guide for the Verifiable Proof-of-Work Engineering Discovery & Hiring Platform. It explains the project’s problem and solution, what the current code actually implements, how to present it clearly, and how to answer difficult questions without overstating the MVP.')
p('Prepared from the current repository: README.md, docs/architecture.md, docs/setup.md, docs/deployment.md, docs/security.md, docs/limitations.md, docs/decision-log.md, ai.md, backend/app.py, backend/seed_demo.py and the React interface.')
h2('The most important message')
p('Present ProofForge as a proof-centered hiring workflow with a working API foundation and a partially synthetic demo interface. Some authentication and assessment/review API flows persist; several role-workspace views still use synthetic browser state. That distinction makes your explanation credible.')
h2('What this guide contains')
for v in ['Plain-language problem and solution','Short and extended jury pitches','A step-by-step product demo','Implemented versus demo versus planned features','Architecture, stack, security and trade-offs','Likely jury questions with answer scripts','A prioritized future roadmap']: bullet(v)

h1('1 Project overview')
p('ProofForge Cyber, presented under the product name FREQUENCY, is a cybersecurity talent discovery and hiring platform built around evidence of practical work. A candidate can complete a technical assessment and submit a structured finding. The API can score assessment sections, preserve findings and review records, and expose sanitized proof through recruiter endpoints after a reviewer decision. The intended outcome is a hiring conversation grounded in inspectable work rather than resume claims alone.')
p('The core problem is a proof gap: a certificate, tool list or resume bullet is a claim about capability, but does not by itself show whether someone can reproduce a security issue, explain its impact, or recommend a fix. Reviewers need a structured way to examine evidence. Recruiters need a readable, privacy-conscious summary of work that has been checked.')
p('The project is not a production cyber range, an automated hiring AI, or an operating-system lockdown tool. It does not currently provision isolated vulnerable containers. The frontend’s role workspaces include synthetic records; the backend has selected persisted flows. Assessment monitoring is a standard browser timer and questionnaire, not camera or screen-share proctoring.')

h1('2 Problem statement and Mysuru context')
h2('Problem statement')
p('Cybersecurity candidates often describe skills through credentials, tool lists and self-reported experience. Hiring teams have limited time to distinguish a candidate who can reproduce and explain a finding from one who can only describe familiar terminology. Candidates with practical ability may also struggle to communicate that ability in a form a reviewer and recruiter can compare.')
p('ProofForge addresses this gap by organizing evidence into a reviewable sequence: candidate work, server-side assessment scoring, human review, evidence-based skill signals, and recruiter discovery of sanitized proof. It aims to reduce uncertainty in screening; it does not eliminate hiring risk or replace interviews.')
h2('Why the Mysuru setting matters')
p('For a Mysuru engineering and hiring audience, the project focuses on making practical cybersecurity ability easier to inspect across candidates, colleges and employers. The repository does not contain verified local labor-market statistics or formal user-study results. Describe the local hiring scenario as context and a hypothesis to validate—not as a measured citywide finding.')
h2('What “solved” looks like')
for v in ['A candidate submits a structured finding and completes a timed baseline assessment.','A reviewer examines the evidence, records a decision and scores the work.','The platform distinguishes API-backed records from demonstration data.','A recruiter reviews candidates passed into the roster by an authorized reviewer and inspects sanitized verified proof.','The evidence supports an interview; a human still makes the hiring decision.']: num(v)

h1('3 Solution explained simply')
p('In one sentence: ProofForge turns candidate security work into structured evidence, asks a human reviewer to verify it, and gives recruiters a sanitized view of the resulting proof.')
h2('Five-step workflow')
for v in [
'Candidate signs in and completes a 40-question assessment arranged as four sections of ten. The overall browser timer is 20 minutes. The UI keeps unfinished question progress in that browser; submitted section scores are calculated by the API.',
'Candidate submits a structured finding with a vulnerability description, evidence, reproduction steps, impact and remediation. The API validates required fields, strips basic HTML tags, hashes evidence with SHA-256 and can flag an exact duplicate for that candidate.',
'Reviewer inspects the submission, chooses verified, rejected or needs changes, and assigns a score. Reviewers may create defense questions; candidate answers and reviewer scores are stored by the API.',
'The API derives capability and proof-confidence signals from persisted findings, assessment results and review-related evidence. These are indicators for review, not a guarantee of job performance or authorship.',
'Reviewer explicitly approves a candidate for the recruiter roster. Recruiter discovery filters the approved roster; recruiter API responses are limited to sanitized verified proof and match reasons. Interview invitations can be persisted by the API.'
]: num(v)
h2('Three roles')
table(['Role','What they need','What the platform supports'],[
('Candidate','A clear way to demonstrate skill and understand feedback.','Assessment, structured findings and technical defense; several workspace actions remain browser-local demo state.'),
('Reviewer','A consistent way to examine proof and record a defensible decision.','Persisted API review decisions, defense questions and scoring; human review remains authoritative.'),
('Recruiter','A safe, quick way to inspect evidence-backed candidates.','API roster requires reviewer approval; verified proof is sanitized. Some recruiter UI records are synthetic.')],[1.1,2.3,3.9])

h1('4 What is implemented today')
p('Use this section to keep claims aligned with the repository. “The design includes” is not the same as “the system does.”')
table(['Area','Supported by current code','Current boundary'],[
('Authentication','Flask API registration and login; scrypt password hashes; JWT access tokens; role checks; logout token revocation; rate limits.','Tokens live in browser session storage for up to four hours. The Render blueprint currently includes fixed demo password values and wildcard CORS; fix and rotate these before real public use.'),
('Assessment','Four sections, ten questions each; API omits answer keys; it scores ten answers and stores one result per candidate/section.','Timer and unfinished question progress are browser-side; no server-side attempt history, randomized order or strong proctoring exists.'),
('Findings and review','API validates structured fields, hashes evidence, flags exact same-candidate duplicates, stores findings, reviewer status/score, defense prompts and rubric data.','A hash detects exact matching content, not authorship. Human review can be mistaken.'),
('Recruiter proof','API roster includes only candidates with reviewer approval; verified proof is sanitized; invitation API routes exist.','Frontend recruiter discovery includes synthetic records and is not fully wired to persisted API data.'),
('Labs','The UI/API lists lab topics and has a synthetic lab flow.','No Dockerized per-candidate vulnerable target is provisioned. Do not claim real isolated labs or authorization to test linked external websites.'),
('Reports','API provides document upload/review endpoints for PDF/Word data and reviewer decisions.','Files are stored as base64 in the database; no malware scanning or mature object-storage pipeline is present.'),
('Monitoring','Camera/screen-share assessment features were removed; hosted permissions policy disables camera, microphone and display capture.','No live proctor feed, kiosk enforcement or OS-level monitoring. The timer is not a secure kiosk.'),
('Reset','Recruiter-only API reset deletes candidate outcomes, findings, scores, documents, approvals, invitations and audit entries.','User accounts are preserved so staff and candidates can sign in again.')],[1.05,3.1,3.15])

h1('5 System architecture')
p('A React and TypeScript single-page frontend communicates with a modular Flask JSON API. SQLAlchemy persists selected API workflows in PostgreSQL in deployment and SQLite locally. The API also writes audit events and computes SHA-256 evidence digests. The frontend has role-based workspaces, but some remain synthetic demonstrations rather than clients of persisted API records.')
h2('Data flow to describe aloud')
for v in ['Candidate UI → Flask API: assessment answers and structured finding submissions.','Reviewer API flow: retrieve submissions, record a decision, ask a defense question and evaluate the answer.','Flask API → PostgreSQL: accounts, findings, assessment section results, review/defense records, approvals, invitations and audit events.','Recruiter API: filters reviewer-approved candidates and returns verified proof in sanitized form.','Candidate UI → authorized lab is the intended flow. Current lab UI is synthetic; managed browser/network policy and isolated lab infrastructure are future deployment layers.']: bullet(v)
h2('Backend records worth knowing')
table(['Record','Why it exists'],[('User','Identity, email, password hash and role.'),('Submission','Structured finding, evidence hash, status and score.'),('AssessmentResult','One score per candidate per section; unique candidate/section pair.'),('DefenseQuestion','Reviewer question, candidate response and evaluation.'),('ReviewerMetrics','Optional report-quality and evidence-consistency scores.'),('RecruiterCandidateApproval','Explicit reviewer decision to pass a candidate into recruiter discovery.'),('Invitation and Audit','Interview invitation state and selected event history.')],[2.0,4.6])

h1('6 Technology stack')
table(['Layer','Technology','Purpose'],[
('Frontend','React 19, TypeScript 7, Vite 8, Lucide React, CSS','Role workspaces, assessment interactions and responsive interface.'),
('API','Python, Flask 3, Flask-SQLAlchemy, Flask-JWT-Extended','JSON endpoints, authentication, role authorization, persistence and scoring.'),
('Data','PostgreSQL hosted; SQLite local; SQLAlchemy ORM','Relational records for selected review and assessment workflows.'),
('Security','Werkzeug scrypt hashes, JWTs, Flask-Limiter, role decorators, SHA-256 evidence digest','Baseline controls; not a full production security program.'),
('Deployment','Render static frontend, Gunicorn API, managed PostgreSQL; local Vite + Flask','Hackathon hosting and development.'),
('AI','Development assistance as disclosed in ai.md','No runtime LLM decides verification, scoring or hiring. Reviewer suggestions are deterministic demo content.')],[1.1,2.75,3.05])

h1('7 Security privacy and responsible claims')
p('The API uses server-side roles and ownership checks for selected endpoints, stores password hashes rather than plaintext passwords, keeps assessment answer keys on the server, and records selected review and audit events. The submission path strips basic HTML tags and enforces a request-size limit. These controls reduce some common risks; they do not make the project production-ready.')
h2('Important security boundaries')
for v in [
'Do not call the platform “unhackable,” “fully secure,” or “proctored.” It is an MVP with basic controls and documented gaps.',
'Never target a website or system without written scope. Links in the interface are not authorization; the HackMysuru site is explicitly view-only unless its operator grants scope.',
'The current Render config sets FRONTEND_ORIGIN to wildcard and contains fixed demo password settings. Before public real use, replace these with managed secrets, rotate exposed values, set the exact frontend origin, and use shared rate-limit storage.',
'Tokens are held in session storage. A same-origin script injection could access them; production should move to Secure, HttpOnly cookies and undergo security review.',
'Uploaded reports are stored as base64 in the database; malware scanning and mature retention/redaction controls are absent. Use synthetic or non-sensitive data in the demo.',
'The webpage cannot control Windows apps, Task Manager, Alt+Tab, other browsers or the operating system. The current assessment has no camera or screen capture.'
]: bullet(v)

h1('8 How to explain it professionally to judges')
h2('Thirty-second pitch')
p('“We built FREQUENCY, a proof-of-work platform for cybersecurity hiring. The problem is that resumes and certificates describe security skills but do not make practical ability easy to verify. ProofForge structures candidate findings and assessment results, gives reviewers the final verification decision, and presents recruiters with sanitized evidence. Our MVP has persisted API flows for authentication, assessment scoring and review, while several role-workspace screens still use synthetic data. Our next step is to connect those screens to the API and provision properly isolated labs.”')
h2('Two-minute pitch')
p('“Cybersecurity hiring has a practical evidence problem. A candidate can list tools or certificates, but an employer still needs to know whether that person can reproduce a finding, explain its impact and propose a useful fix. That uncertainty affects candidates, reviewers and recruiters.')
p('“FREQUENCY—ProofForge Cyber—organizes that evidence in five stages. Candidates complete a timed baseline assessment and submit structured findings. The API keeps the answer key server-side, scores each section and hashes submitted evidence. Reviewers examine the finding, record a decision and can ask technical defense questions. Only candidates explicitly approved by a reviewer enter recruiter discovery, where verified proof is returned in sanitized form.')
p('“We chose deterministic scoring and human review over an LLM as the hiring authority. That creates reviewer delay, but it keeps decisions inspectable and accountable. We are careful about the current boundary: authentication and selected API workflows persist, while parts of the candidate and recruiter interface still show synthetic demo state. We have not provisioned isolated vulnerable containers, and our assessment is not camera-proctored. Our immediate roadmap is to connect the remaining role pages, add isolated egress-restricted labs, strengthen evidence storage and test real peak loads.”')
h2('Present claims carefully')
for v in ['Say “evidence-informed screening,” not “guaranteed hiring accuracy.”','Say “human-verified finding” only when a reviewer actually marked that record verified.','Say “the API supports persisted review flow,” not “every screen is connected to the database.”','Say “synthetic lab interface,” not “isolated cyber range.”','Say “browser timer,” not “tamper-proof proctoring.”','Say “we estimate 167 requests per second under an even-spread assumption,” not “we proved the system handles that traffic.”']: bullet(v)

h1('9 Suggested live demo')
p('Keep the demonstration to about four minutes. Narrate what each action proves. Before presenting, confirm that the app is running, API health responds, demo accounts are available, and no real personal data is visible.')
for v in [
'Start with the three roles: candidate, reviewer and recruiter. Explain that the account record supplies the role; a user-controlled workspace picker does not grant privileges.',
'Candidate: open the assessment instructions. Explain 40 questions, four sections, a 20-minute overall browser timer, and server-side answer keys.',
'Show the structured finding form or a prepared sample. Point out evidence, reproduction, impact and remediation—not just the vulnerability label.',
'Reviewer: show a pending finding and the decision choices. If available, demonstrate a defense question and recorded evaluation.',
'Recruiter: show the approved-candidate roster and sanitized proof. Explain that the API gates the roster on reviewer approval while some visible frontend cards may be synthetic.',
'Close with one limitation and one next step: connect remaining role pages to persisted records, then provision isolated labs with network egress restrictions.'
]: num(v)
h2('If the live demo fails')
p('Do not hide the failure or improvise a claim. Say: “The hosted demo is not responding right now. I can still explain the API workflow and show the architecture and implementation. The app uses Render for the frontend and API, with PostgreSQL in the hosted blueprint.” Then walk through the deck and identify which screenshots or records are actual captures.')

h1('10 Jury questions and model answers')
p('Use these as speaking notes, not word-for-word memorized claims. Answer directly first, explain the trade-off second, and identify planned work as planned.')
qa('1 What exact problem are you solving?','We address the gap between cybersecurity skill claims and reviewable proof. The product helps candidates submit structured work, reviewers assess it, and recruiters inspect sanitized verified evidence. It supports screening; it does not replace a technical interview.')
qa('2 Why is this relevant to Mysuru?','We framed it for the Mysuru engineering and hiring ecosystem, where candidates, reviewers and employers need a shared way to discuss practical security ability. We do not have a measured citywide study in this repository, so we present that as context to validate with local users, not as a statistic.')
qa('3 Who are your users?','Candidates demonstrate work; reviewers examine and score evidence; recruiters discover candidates explicitly approved. Each role has different permissions and information needs.')
qa('4 What is your core differentiator?','The product connects structured evidence to human review and recruiter-facing proof. Claimed tools do not count as verified skill, and a reviewer—not an AI model—controls the final finding decision.')
qa('5 How do you know a finding is real?','The API records evidence and reproduction information, flags exact duplicate evidence for a candidate, and waits for reviewer decision. SHA-256 is an integrity/deduplication signal, not proof of authorship; the reviewer judges reproducibility and context.')
qa('6 What does Security DNA measure?','It presents capability and proof-confidence signals derived from selected assessment and reviewed evidence. It is a product metric to organize evidence, not a validated predictor of job performance. Some frontend profiles remain synthetic.')
qa('7 Why separate capability and proof confidence?','A skill estimate and the amount or quality of supporting evidence are different questions. A claim without reviewed proof should not look identical to a signal supported by reviewed work. The score still needs validation and transparent calibration before real hiring use.')
qa('8 Why not use an LLM to grade submissions?','An LLM could speed up summaries, but using it as final grader would make scores harder to reproduce and defend, especially when evidence itself is adversarial input. We chose human decisions and deterministic API scoring. The displayed reviewer suggestion is demo content, not a runtime model.')
qa('9 Does the product use AI?','AI tools assisted development as documented in ai.md. Runtime has no LLM verifying findings, ranking candidates or deciding hiring. Any future AI feature should be disclosed and advisory, with human accountability.')
qa('10 Are your labs isolated and safe to attack?','Not yet. The current lab experience is synthetic and does not provision per-candidate Docker containers. Interface links are not authorization to test external websites. Real labs need isolated ephemeral environments, egress denial, reset controls and explicit scope.')
qa('11 Is the assessment proctored?','No. It is a timed browser questionnaire. Camera and screen-share functions are not part of the current product, and the webpage does not control the operating system or other apps.')
qa('12 Can candidates cheat on the assessment?','The API keeps correct answer indexes server-side and computes submitted section scores. Per-question attempt history, randomized order and server-authoritative timing are not implemented, so we call it a baseline assessment—not a cheat-proof exam.')
qa('13 How does your timer work?','The frontend starts an overall 20-minute timer when the candidate begins and stores progress in that browser so it continues through refreshes. It is not a trusted server-side clock. Production should store start/end times and enforce timing on the server.')
qa('14 What security controls do you have?','Passwords are scrypt-hashed; API routes use JWT authentication and role checks; selected routes enforce ownership; rate limits apply; evidence is hashed; selected actions are audited. These are baseline controls, not a security certification.')
qa('15 Is it production-ready?','No. It is a hackathon MVP. Render config needs private demo credentials, exact CORS origins and shared rate-limit storage; frontend flows remain synthetic; evidence storage/retention need work; isolated labs are not provisioned.')
qa('16 How is candidate privacy handled?','Recruiter proof routes sanitize verified records and avoid exposing raw evidence in the recruiter summary. But uploads are stored as base64 in the database, redaction is pattern-based, and retention controls are incomplete. We use synthetic data in the demo and need a reviewed storage/access/deletion policy before real reports.')
qa('17 Can a recruiter see every candidate?','The API recruiter list is based on candidates explicitly approved by a reviewer. The recruiter proof endpoint checks that approval too. Frontend demo cards can be synthetic, so distinguish visible sample data from persisted API behavior.')
qa('18 How do you prevent a candidate from becoming a reviewer?','Public registration is candidate-only. Reviewer and recruiter roles must be provisioned through a trusted administrative process; API role checks protect staff actions.')
qa('19 What is persisted and what is local-only?','The API persists accounts, findings, assessment section results, reviewer decisions, defense records, document review metadata, approvals, invitations and audit events. Unfinished question progress is stored in the browser; several workspace actions and profile cards are synthetic local demo state.')
qa('20 What happens when recruiter presses reset?','The recruiter-only reset API deletes candidate outcomes and workflow records, including findings, scores, reviews, approvals, invitations and audit entries. User accounts remain so the demo can be used again. Verify this behavior before the demo and explain it cannot be undone.')
qa('21 How do you scale to all of Mysuru?','We have not load-tested that scale. As an estimate, 50,000 candidates times four sections is 200,000 score submissions. Spread evenly across 20 minutes, that is about 167 requests per second; bursts could be higher. Synchronous scoring, separate audit commits, one configured worker and in-memory rate limits are not validated for that load.')
qa('22 What would you change first for scale?','Add durable, idempotent submission ingestion with an outbox/worker path for scoring and audit. Move rate limits to shared Redis, then load-test representative bursts and tune workers and database capacity. The 167 requests-per-second figure is arithmetic, not a benchmark.')
qa('23 What happens with poor internet or offline use?','The app is a lightweight web client and unfinished assessment progress is stored in that browser. It has no complete offline synchronization guarantee; API-backed actions need connectivity. A future enhancement is an offline queue with conflict handling and clear sync status.')
qa('24 How will you support Kannada and different devices?','Responsive web UI is the base, but the repository does not prove a Kannada-first experience or validated accessibility with local users. Test Kannada/English terms, low bandwidth, small screens, keyboard access and assistive technology before claiming those needs are solved.')
qa('25 How do you avoid reviewer bias?','The system records human decisions and offers rubric dimensions, but a calibrated reviewer program is not complete. Add anchored scoring examples, reviewer calibration, disagreement monitoring, candidate appeals and periodic fairness audits.')
qa('26 How do you know this improves hiring outcomes?','We do not have hiring outcome evidence yet. A pilot should measure reviewer time and agreement, candidate completion, finding decisions, recruiter usefulness and interview progression, while checking for unfair impact.')
qa('27 What is the biggest trade-off you accepted?','Human verification takes longer and makes reviewer capacity a bottleneck. We prefer a decision a reviewer can explain over a faster opaque automated score. We also accept that current UI/API integration is incomplete and disclose it.')
qa('28 What is your business model?','The repository implements a prototype, not a validated business model. A plausible direction is institution or employer subscriptions for assessment and review workspaces, but pricing and willingness to pay need interviews and pilot evidence before claiming product-market fit.')
qa('29 What did the team build versus AI?','The codebase contains the React/TypeScript client, Flask API, data models and deployment configuration. ai.md discloses AI assistance for scaffolding, question pools and documentation. No runtime AI makes verification or hiring decisions. State each member’s actual contribution if asked; do not invent it.')
qa('30 What are your next three steps?','Connect remaining role pages to authenticated persisted data; provision isolated resettable labs with egress blocked; then strengthen operations and security through managed secrets, exact CORS origins, shared rate limits, secure evidence storage, retention, monitoring and load tests.')
qa('31 What limitation should you admit first?','Several role views still use synthetic local data. The API has persisted workflows, but the product is not fully integrated end to end. We call that out and make UI/API integration the first roadmap item.')
qa('32 What if the reviewer is wrong?','A reviewer can make a mistake. The system should support a second review or appeal, transparent rationale and reviewer calibration. Those governance processes are not complete in the MVP.')

h1('11 Future enhancements')
h2('Priority one connect the product end to end')
p('Replace synthetic candidate, reviewer and recruiter screens with authenticated API data. Add visible sync status, error handling and audit context. Make every result traceable to its source record and label demo data until integration is complete.')
h2('Priority two provision safe practical labs')
p('Run a disposable isolated container per candidate/session, deny outbound network by default, apply CPU/memory/time quotas, support reset and capture only necessary outputs. Publish explicit authorized scope. Managed kiosk and proxy controls must be organization-managed; a normal webpage cannot enforce OS-level restrictions.')
h2('Priority three improve assessment integrity')
p('Move timing enforcement to the server, store start/submission times, preserve attempt history, randomize equivalent question sets, and support accommodations and appeals. Surveillance is not a substitute for assessment design.')
h2('Priority four harden sensitive evidence handling')
p('Move files from the relational database into private object storage, validate file type and size, scan uploads, encrypt data, apply role-scoped access, redact sensitive content with reviewer confirmation, and define retention/deletion rules.')
h2('Priority five scale and operate responsibly')
p('Use managed secrets and remove committed demo passwords, set precise CORS origins, move rate limits to shared Redis, add database migrations/backups, alerting, application logs and load tests. Use a scoring/audit queue only after defining idempotency and consistency.')
h2('Priority six validate usefulness and fairness')
p('Pilot with candidates, reviewers and recruiters. Measure completion, reviewer agreement, review time, false-positive concerns and recruiter usefulness. Test language, device constraints, accessibility, bias and appeals with real participants.')

h1('12 Local setup and demo preparation')
p('Use docs/setup.md as the setup source of truth. In brief: from the repository root run npm install and npm run dev; in backend create/activate a Python virtual environment, install backend/requirements.txt and start the Flask app. To seed demo accounts, set a private password of at least 12 characters before running backend/seed_demo.py. Do not present a committed/default password as secure.')
h2('Seeded demo sign-ins')
p('The seed script creates candidate accounts Ananya Rao (ananya.demo@example.invalid), Rohan Mehta (rohan.demo@example.invalid) and Maya Iyer (maya.demo@example.invalid), plus reviewer Samira Khan and recruiter Jordan Davis. Keep credentials private; set environment variables locally and never put password values in a deck, screenshot or public document.')
h2('Before jury day')
for v in ['Confirm frontend and API start and the API health endpoint responds.','Confirm candidate, reviewer and recruiter sign-in works with environment-configured demo passwords.','Prepare one synthetic finding and know its reviewer status.','Know which interface records are synthetic and which API records persist.','Use safe sample evidence, not personal or sensitive reports.','Keep the short pitch and first limitation ready; do not claim isolated labs or proctoring.','If reset is used, verify that accounts remain and the workflow starts cleanly.']: bullet(v)

h1('13 Decision summary and scale calculation')
p('Chosen approach: structured evidence, deterministic server-side assessment scoring and human reviewer authority. Rejected approach: an LLM as final vulnerability grader. An LLM might provide faster feedback, but a final score based on adversarial evidence would be harder to reproduce and defend. The accepted cost is human-review delay and a reviewer-capacity bottleneck.')
h2('Transparent scale estimate')
p('The API stores one assessment result per candidate and section. With 50,000 candidates completing four sections, the estimate is 50,000 × 4 = 200,000 score submissions. Spread evenly across 20 minutes (1,200 seconds), the average is 200,000 ÷ 1,200 ≈ 167 submissions per second. The real peak may be higher if many candidates submit at one deadline. This is arithmetic from the data model, not a benchmark or current capacity claim.')
p('A first engineering change would be durable, idempotent submission ingestion with an outbox/worker path for scoring and audit. Shared Redis rate limits and load tests should precede horizontal scaling. The queue adds infrastructure and asynchronous result handling, which must be reflected in the candidate UI.')

doc.add_page_break()
h1('14 Useful terms')
table(['Term','Plain-English explanation'],[
('Proof of work','Evidence of a completed task, rather than a claim or credential.'),('Finding','A structured security issue with evidence, reproduction, impact and remediation.'),('Reviewer verification','A human decision about whether a finding is sufficiently supported.'),('SHA-256 digest','A fixed-length hash used to flag identical evidence strings; it does not prove originality.'),('Capability','A signal about demonstrated skill; not a guarantee of job performance.'),('Proof confidence','A separate signal about the amount/quality of evidence for a capability.'),('JWT','A signed access token used by the API to identify an authenticated account.'),('RBAC','Role-based access control: permissions depend on candidate/reviewer/recruiter role.'),('Egress restriction','A network rule that blocks a lab from reaching unauthorized systems.'),('Synthetic UI data','Sample browser-side records used to demonstrate a screen; not candidate history.')],[1.55,5.05])

h1('15 Repository reference')
for label,path in [('Main project overview','README.md'),('Architecture','docs/architecture.md'),('Local setup','docs/setup.md'),('Deployment','docs/deployment.md'),('Security notes','docs/security.md'),('Limitations and roadmap','docs/limitations.md'),('Decision rationale','docs/decision-log.md'),('AI disclosure','ai.md'),('API implementation','backend/app.py'),('Demo seed','backend/seed_demo.py'),('React interface','src/ui.tsx')]:
 x=doc.add_paragraph(style='List Bullet'); x.paragraph_format.space_after=Pt(2); r=x.add_run(label+': '); r.bold=True; x.add_run(path)
p('Use this guide together with the live application and source code. If the interface changes, update the demo script and speaking notes before the next jury presentation.')

doc.core_properties.title='ProofForge Cyber Jury Guide'
doc.core_properties.subject='Project explanation, jury pitch, demo walkthrough, questions and roadmap'
doc.core_properties.author='Team HM26-C3B5'; doc.core_properties.keywords='ProofForge, FREQUENCY, cybersecurity hiring, jury guide'
doc.save(OUT); print(OUT)

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_THEME_COLOR

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / 'output' / 'pptx' / 'HM26-C3B5_presentation.pptx'
DEST.parent.mkdir(parents=True, exist_ok=True)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
COLORS = {'paper':'F7F7F2','green':'24563E','dark':'192D25','mint':'EDF3EB','ink':'25332C','muted':'65756C','line':'DCE4D9','warm':'F1EEE4','orange':'C56D3E','white':'FFFFFF'}

def rgb(key):
    h=COLORS.get(key,key).lstrip('#'); return RGBColor.from_string(h)

def rect(slide,x,y,w,h,fill='white',line=None,radius=True):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    sh.line.color.rgb=rgb(line or fill); sh.line.width=Pt(0.8)
    if radius:
        try: sh.adjustments[0]=0.09
        except Exception: pass
    return sh

def text(slide,txt,x,y,w,h,size=18,color='ink',bold=False,font='Aptos',align=PP_ALIGN.LEFT,margin=0.03,valign=MSO_ANCHOR.TOP):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(margin); tf.margin_right=Inches(margin); tf.margin_top=Inches(margin); tf.margin_bottom=Inches(margin)
    tf.vertical_anchor=valign
    for i,line in enumerate(txt.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.alignment=align
        p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color)
        p.space_after=Pt(5)
    return box

def line(slide,x1,y1,x2,y2,color='green',width=1.2):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=rgb(color); sh.line.width=Pt(width); sh.line.end_arrowhead=True
    return sh

def base(n,title):
    s=prs.slides.add_slide(BLANK); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb('paper')
    rect(s,0,0,13.333,0.12,'green',radius=False)
    text(s,'FREQUENCY  /  PROOFFORGE CYBER',0.72,0.32,6.3,0.24,10,'muted',True)
    text(s,f'{n:02d} / 10',11.7,0.32,0.9,0.24,10,'muted',True,align=PP_ALIGN.RIGHT)
    text(s,title,0.72,0.78,11.9,0.55,28,'dark',True)
    rect(s,0.72,1.47,11.9,0.012,'line',radius=False)
    rect(s,0.72,7.03,0.06,0.3,'green',radius=False)
    text(s,'TEAM HM26-C3B5  ·  EVIDENCE BEFORE CLAIMS',0.92,7.05,6.5,0.18,8,'muted')
    text(s,'FREQUENCY',11.2,7.05,1.4,0.18,8,'muted',align=PP_ALIGN.RIGHT)
    return s

def add_slide1():
    s=base(1,'FREQUENCY')
    rect(s,0.72,2.0,0.1,3.3,'green',radius=False)
    text(s,'VERIFIABLE PROOF-OF-WORK\nENGINEERING DISCOVERY & HIRING',1.05,1.95,10.9,1.35,34,'dark',True)
    text(s,'Proof you can inspect. Skills you can trust.',1.07,3.54,10.7,0.45,20,'green',True)
    text(s,'Team HM26-C3B5  ·  Members: [add names]  ·  Colleges: [add colleges]',1.07,4.2,10.9,0.45,17,'ink')
    text(s,'Selected problem · The proof gap in cybersecurity hiring',1.07,4.8,10.9,0.4,16,'muted')

def add_slide2():
    s=base(2,'THE PROOF GAP')
    text(s,'Resumes and certificates show claims. They rarely show reproducible work, sound reasoning, or how a candidate explains a finding.',0.78,1.95,11.6,1.45,25,'dark',True)
    rect(s,0.75,4.0,11.85,1.9,'mint')
    text(s,'A MYSURU HIRING SCENARIO',1.0,4.28,10.9,0.3,15,'green',True)
    text(s,'A local team screens many applicants. A capable candidate without a familiar credential is easy to miss; an inflated claim is hard to verify.',1.0,4.72,10.9,0.72,18,'ink')
    text(s,'Candidates lose visibility. Reviewers spend time rechecking. Recruiters lack comparable evidence.',1.0,5.55,10.9,0.32,15,'muted')

def add_slide3():
    s=base(3,'WHO IT SERVES')
    cards=[('CANDIDATE','Demonstrates work through a finding, assessment and technical defense. Needs a clear path and low-bandwidth access.'),('REVIEWER','Checks reproduction evidence, records a decision and scores defense. Human judgment remains authoritative.'),('RECRUITER','Filters evidence-backed profiles and reviews sanitized proof before inviting candidates.')]
    for i,(head,body) in enumerate(cards):
        x=0.72+i*4.08; rect(s,x,2.1,3.72,3.55,'white','line')
        text(s,head,x+0.25,2.45,3.2,0.35,16,'green',True)
        text(s,body,x+0.25,3.0,3.2,2.2,18,'ink')

def add_slide4():
    s=base(4,'FROM WORK TO PROOF')
    steps=['Candidate completes a section assessment and submits a structured finding.','API scores answers server-side and hashes submitted evidence.','Reviewer reproduces, questions and records a decision.','Verified records inform capability and proof-confidence signals.','Recruiter inspects sanitized proof and can invite the candidate.']
    for i,t in enumerate(steps):
        y=1.95+i*0.72
        circ=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(0.78),Inches(y),Inches(0.42),Inches(0.42)); circ.fill.solid(); circ.fill.fore_color.rgb=rgb('mint'); circ.line.color.rgb=rgb('mint')
        text(s,f'{i+1:02d}',0.78,y+0.09,0.42,0.2,11,'green',True,align=PP_ALIGN.CENTER)
        text(s,t,1.38,y+0.02,10.9,0.4,17,'ink')
    rect(s,0.75,5.82,11.85,0.78,'warm')
    text(s,'PRODUCT CAPTURE NEEDED',0.98,5.98,2.4,0.25,13,'orange',True)
    text(s,'No genuine screenshots are checked in. Replace this with a capture of the running product.',3.28,5.97,8.9,0.33,15,'muted')

def add_slide5():
    s=base(5,'ARCHITECTURE')
    text(s,'React + TypeScript role workspaces  →  Flask modular API  →  PostgreSQL',0.78,1.7,11.8,0.38,17,'green',True)
    boxes=[('Candidate UI',0.85,2.35,2.0),('Reviewer UI',4.35,2.35,2.0),('Recruiter UI',7.85,2.35,2.0),('Flask modular API',4.35,3.65,2.0),('PostgreSQL',4.35,5.05,2.0),('Audit events',8.8,3.65,1.9),('Evidence SHA-256',8.8,5.05,1.9),('Authorized isolated lab',0.85,5.05,2.0),('Managed browser / network allowlist',0.72,3.65,2.35)]
    centers={}
    for name,x,y,w in boxes:
        rect(s,x,y,w,0.62,'mint' if 'UI' in name else 'white','line')
        text(s,name,x+0.08,y+0.18,w-0.16,0.28,11,'dark',True,align=PP_ALIGN.CENTER)
        centers[name]=(x+w/2,y+0.31)
    api=centers['Flask modular API']
    for name in ['Candidate UI','Reviewer UI','Recruiter UI']:
        x,y=centers[name]; line(s,x,y+0.31,api[0],3.65)
    line(s,api[0],4.27,api[0],5.05); line(s,6.35,3.96,8.8,3.96); line(s,6.35,4.1,8.8,5.25)
    # Browser/network policy gates access to the authorized lab, as in docs/architecture.md.
    line(s,1.9,4.27,1.9,5.05)
    text(s,'API records audit events and hashes evidence. Candidate lab UI is synthetic; no isolated target is provisioned.',0.78,6.0,11.8,0.58,12,'muted')

def add_slide6():
    s=base(6,'CONSTRAINTS & HONEST STATUS')
    rows=['Evidence integrity · SHA-256 hash and duplicate flag in API  |  PARTIAL','Jurisdiction · No isolated lab provisioned; do not test arbitrary targets  |  NOT BUILT','Priority · Human reviewer controls verification  |  HANDLED','Bad input · API validates assessment shape and finding fields  |  PARTIAL','Offline · Local SQLite fallback; no offline sync guarantee  |  PARTIAL']
    for i,t in enumerate(rows):
        y=1.95+i*0.87; rect(s,0.75,y,11.85,0.62,'white','line'); text(s,t,1.0,y+0.16,11.3,0.31,17,'ink')

def add_slide7():
    s=base(7,'THE KEY DECISION')
    panels=[('CHOSEN','Deterministic server-side scoring + structured evidence + human reviewer authority','mint','green'),('REJECTED','LLM as final grader: faster feedback, but less repeatable and harder to defend on adversarial evidence.','warm','orange')]
    for i,(head,body,bg,accent) in enumerate(panels):
        x=0.75+i*6.1; rect(s,x,2.0,5.75,2.65,bg,bg)
        text(s,head,x+0.3,2.3,5.1,0.35,16,accent,True)
        text(s,body,x+0.3,2.9,5.1,1.45,19,'ink')
    text(s,'TRADE-OFF ACCEPTED',0.78,5.15,4.0,0.3,16,'dark',True)
    text(s,'Human review takes longer and creates a reviewer-capacity bottleneck. The MVP favors explainable decisions.',0.78,5.62,11.7,0.7,17,'muted')

def add_slide8():
    s=base(8,'WHAT THE CODE PROVES')
    cards=[('API flow','Assessment answer keys stay server-side; each section score is persisted.'),('Review flow','Human reviewer decisions, defense responses and rubric values have API persistence.'),('Discovery boundary','API sanitizes verified proof. Some candidate and recruiter screens still use synthetic data.')]
    for i,(head,body) in enumerate(cards):
        x=0.72+i*4.08; rect(s,x,2.0,3.78,2.0,'white','line')
        text(s,head,x+0.22,2.27,3.35,0.3,15,'green',True)
        text(s,body,x+0.22,2.8,3.35,1.0,17,'ink')
    rect(s,0.75,4.5,11.85,1.3,'warm')
    text(s,'REAL CAPTURE SLOTS',0.98,4.76,3.0,0.3,14,'orange',True)
    text(s,'Add 2–3 genuine screenshots: candidate assessment, reviewer decision, recruiter sanitized proof. None are checked in.',0.98,5.15,11.1,0.5,16,'muted')

def add_slide9():
    s=base(9,'LIMITS & SCALE')
    text(s,'50,000 candidates × 4 sections = 200,000 score submissions.',0.78,1.85,11.8,0.5,23,'green',True)
    rows=['Over 20 minutes: ~167 requests/second average; deadline bursts are higher.','First pressure point: synchronous score + separate audit commits; one configured worker; in-memory rate limit.','First change: durable idempotent submission/outbox queue; shared Redis limits; then load-test.','Other gaps: synthetic UI data; no isolated labs; no secure file-upload/redaction pipeline.']
    for i,t in enumerate(rows):
        y=2.7+i*0.78; text(s,'•',0.85,y,0.25,0.36,20,'green',True); text(s,t,1.2,y,11.0,0.52,17,'ink')

def add_slide10():
    s=base(10,'NEXT & DISCLOSURE')
    todos=['1  Connect role pages to authenticated persisted records.','2  Provision isolated, resettable labs with egress denied.','3  Add shared rate limiting, monitoring and peak-load tests.','AI disclosure: Claude/Gemini assisted development; no runtime LLM makes hiring decisions.']
    for i,t in enumerate(todos): text(s,t,0.82,1.88+i*0.66,11.7,0.45,17,'ink')
    rect(s,0.75,4.95,11.85,1.5,'mint')
    text(s,'LIVE MVP',1.0,5.2,1.5,0.25,13,'green',True); text(s,'https://proofforge-cyber.onrender.com',2.55,5.16,9.7,0.32,15,'green')
    text(s,'REPOSITORY',1.0,5.65,1.5,0.25,13,'green',True); text(s,'github.com/yogesh-37911/HM26-C3B5-Phase2',2.55,5.61,9.7,0.32,15,'ink')
    text(s,'Team members / colleges: [add names]',1.0,6.05,10.8,0.25,14,'muted')

for fn in [add_slide1,add_slide2,add_slide3,add_slide4,add_slide5,add_slide6,add_slide7,add_slide8,add_slide9,add_slide10]: fn()
prs.core_properties.title='HM26-C3B5 Presentation — FREQUENCY / ProofForge Cyber'
prs.core_properties.author='Team HM26-C3B5'
prs.core_properties.subject='Verifiable Proof-of-Work Engineering Discovery & Hiring Platform'
prs.save(DEST)
print(DEST)
